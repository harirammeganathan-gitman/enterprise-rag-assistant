"""Load supported file formats while preserving page-level provenance."""

from hashlib import sha256
from pathlib import Path

from docx import Document
from pypdf import PdfReader


def load_document(path: Path) -> tuple[list[tuple[str, int | None]], str]:
    """Return extracted page/text units and a content hash for one file."""
    suffix = path.suffix.lower()
    content_hash = sha256(path.read_bytes()).hexdigest()

    if suffix == ".pdf":
        pages = [
            (page.extract_text() or "", page_number)
            for page_number, page in enumerate(PdfReader(path).pages, start=1)
        ]
    elif suffix == ".docx":
        document = Document(path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        tables = [
            " ".join(cell.text for row in table.rows for cell in row.cells)
            for table in document.tables
        ]
        text = "\n".join(part for part in [*paragraphs, *tables] if part.strip())
        pages = [(text, None)]
    elif suffix in {".md", ".markdown"}:
        pages = [(path.read_text(encoding="utf-8"), None)]
    else:
        raise ValueError(f"Unsupported file type: {suffix or '(no extension)'}")

    if not any(text.strip() for text, _ in pages):
        raise ValueError(f"No extractable text found in {path.name}")
    return pages, content_hash
